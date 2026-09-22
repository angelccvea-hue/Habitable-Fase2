"""Importación de plantillas Excel (informe vertical o lote) hacia CasoRojo."""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from django.db import models
from openpyxl import load_workbook

from inspecciones.excel_plantilla import EXPORT_COLUMNS, PRECARGA_FIELDS, ROWS_DEF
from inspecciones.models import CasoRojo, LineaMetrado, PartidaCatalogo

# Campos que no se sobrescriben desde Excel (asignación / auditoría / GPS ORM).
SKIP_IMPORT_FIELDS = frozenset(
    {
        "inspector_asignado",
        "revisor_asignado",
        "coordinador_asignado",
        "created_at",
        "updated_at",
        "lat",
        "lng",
        "id",
        "pk",
        # §11 solo en ficha web (aunque un Excel viejo traiga estas filas)
        "n_fotos",
        "firmas",
    }
)

LABEL_TO_FIELD: dict[str, str] = {
    str(CasoRojo._meta.get_field(fname).verbose_name): fname
    for _sec, etiqueta, fname, _lista, _pre in ROWS_DEF
    if fname and etiqueta
}


class ExcelImportError(Exception):
    """Error de formato o de negocio al importar Excel."""


def _norm(text: Any) -> str:
    if text is None:
        return ""
    return str(text).strip()


def _parse_date(val: Any) -> date | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = _norm(val)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    raise ExcelImportError(f"Fecha no válida: {val!r}")


def _coerce_field(field: models.Field, raw: Any) -> Any:
    if raw is None or raw == "":
        if isinstance(field, (models.CharField, models.TextField)):
            return ""
        return None
    if isinstance(field, models.DateField):
        return _parse_date(raw)
    if isinstance(field, models.DecimalField):
        try:
            return Decimal(str(raw).replace(",", "."))
        except (InvalidOperation, ValueError) as exc:
            raise ExcelImportError(f"{field.name}: decimal inválido ({raw!r})") from exc
    if isinstance(field, (models.IntegerField, models.PositiveIntegerField, models.PositiveSmallIntegerField)):
        try:
            return int(float(str(raw).replace(",", ".")))
        except (TypeError, ValueError) as exc:
            raise ExcelImportError(f"{field.name}: entero inválido ({raw!r})") from exc
    if isinstance(field, models.BooleanField):
        s = _norm(raw).lower()
        if s in {"1", "true", "sí", "si", "yes"}:
            return True
        if s in {"0", "false", "no"}:
            return False
        return bool(raw)
    # TextChoices: valor = etiqueta en este proyecto
    text = _norm(raw)
    if getattr(field, "choices", None):
        valid = {str(c[0]) for c in field.choices}
        labels = {str(c[1]): str(c[0]) for c in field.choices}
        if text in valid:
            return text
        if text in labels:
            return labels[text]
        # Prefijo corto p.ej. "D3" → primera opción que empiece igual
        for choice_val in valid:
            if choice_val.startswith(text) or text.startswith(choice_val.split("—")[0].strip()):
                if len(text) >= 2:
                    matches = [v for v in valid if v.startswith(text) or text in v]
                    if len(matches) == 1:
                        return matches[0]
        # No tumbar todo el Excel por un desplegable mal escrito:
        # el llamador puede saltar el campo con advertencia.
        raise ExcelImportError(f"{field.name}: valor no permitido ({text!r})")
    return text


def _apply_dict_to_caso(
    caso: CasoRojo,
    data: dict[str, Any],
    *,
    overwrite_precarga: bool = False,
) -> tuple[list[str], list[str]]:
    changed: list[str] = []
    warnings: list[str] = []
    for fname, raw in data.items():
        if fname in SKIP_IMPORT_FIELDS or fname not in EXPORT_COLUMNS:
            continue
        if fname == "hab_id":
            continue
        if fname in PRECARGA_FIELDS and not overwrite_precarga:
            # Solo actualiza precarga si viene valor y el campo actual está vacío
            cur = getattr(caso, fname, None)
            if cur not in (None, ""):
                continue
        try:
            field = CasoRojo._meta.get_field(fname)
        except Exception:
            continue
        if not isinstance(field, models.Field) or field.many_to_many or field.one_to_many:
            continue
        # No borrar datos ya cargados con celdas vacías de la plantilla.
        # (Reimportar un Excel parcial no debe limpiar secciones ya llenadas en la ficha.)
        if raw is None or (isinstance(raw, str) and not str(raw).strip()):
            cur = getattr(caso, fname, None)
            if cur not in (None, ""):
                continue
        try:
            new_val = _coerce_field(field, raw)
        except ExcelImportError as exc:
            # Valores de lista inválidos: omitir campo y avisar (no abortar el resto).
            msg = str(exc)
            if "valor no permitido" in msg or "decimal inválido" in msg or "entero inválido" in msg or "Fecha no válida" in msg:
                warnings.append(f"{caso.hab_id}/{fname}: {msg} — campo omitido")
                continue
            raise
        old_val = getattr(caso, fname)
        if old_val != new_val:
            setattr(caso, fname, new_val)
            changed.append(fname)
    # Si el Excel trae corr_*, aplicar a precarga (mismo criterio que el admin)
    if caso.correcciones_estructuradas_llenas():
        aplicados = caso.aplicar_correcciones_a_precarga()
        if aplicados:
            changed.append("precarga_via_correccion")
    return changed, warnings


def parse_informe_sheet(ws) -> dict[str, Any]:
    """Hoja vertical: col A etiqueta / sección, col B valor."""
    data: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        label, value = (row + (None, None))[:2]
        label_s = _norm(label)
        if not label_s:
            continue
        # Encabezados de sección / portada / ayudas
        if value is None and label_s.startswith(
            (
                "1 —",
                "2 —",
                "3 —",
                "4 —",
                "5 —",
                "6 —",
                "7 —",
                "8 —",
                "9 —",
                "10 —",
                "11 —",
                "12 —",
                "INFORME",
            )
        ):
            continue
        if label_s.lower() in {"campo", "valor (rellene aquí)", "cómo llenar", "como llenar"}:
            continue

        # Primero mapear a campo del modelo; no descartar etiquetas válidas
        # (p. ej. «Uso») por heurísticas de la Portada.
        fname = LABEL_TO_FIELD.get(label_s)
        if not fname and label_s in EXPORT_COLUMNS:
            fname = label_s
        if not fname:
            # Filas de Portada / instrucciones (sin campo asociado)
            low = label_s.lower()
            if low.startswith(
                (
                    "azul claro",
                    "plantilla",
                    "edificio",
                    "uso",
                    "devolución",
                    "formato",
                    "casos",
                    "pestaña a rellenar",
                    "▶",
                )
            ):
                continue
            continue
        data[fname] = value
    if "hab_id" not in data:
        # Intentar desde título "INFORME — ID 123"
        title = _norm(ws["A1"].value)
        if "ID" in title:
            import re

            m = re.search(r"ID\s+(\d+)", title, re.I)
            if m:
                data["hab_id"] = int(m.group(1))
    return data


def parse_lote_sheet(ws) -> list[dict[str, Any]]:
    """Hoja Lote_informes: fila 1 claves, fila 2 etiquetas, datos desde fila 3."""
    headers: list[str] = []
    for col in range(1, ws.max_column + 1):
        key = _norm(ws.cell(1, col).value)
        if key:
            headers.append(key)
        else:
            break
    if not headers:
        raise ExcelImportError("Hoja de lote sin encabezados (fila 1).")
    rows: list[dict[str, Any]] = []
    for r in range(3, ws.max_row + 1):
        first = ws.cell(r, 1).value
        if first is None or _norm(first) == "":
            # fila vacía
            if all(ws.cell(r, c).value in (None, "") for c in range(1, len(headers) + 1)):
                continue
        data: dict[str, Any] = {}
        for c, key in enumerate(headers, start=1):
            data[key] = ws.cell(r, c).value
        if data.get("hab_id") in (None, ""):
            continue
        rows.append(data)
    return rows


# Hojas auxiliares de la plantilla (nunca son el informe vertical)
_SKIP_SHEETS = frozenset(
    {
        "Portada",
        "Listas_desplegables",
        "Lote_informes",
        "Metrados",
        "Procedimientos_sel",
        "Catalogo_partidas",
        "Catalogo_procedimientos",
    }
)


def _sheet_hab_id_from_name(name: str) -> int | None:
    """Pestañas tipo «197189_Residencia_M» → 197189."""
    import re

    m = re.match(r"^(\d{4,})(?:_|$)", (name or "").strip())
    return int(m.group(1)) if m else None


def _looks_like_informe_sheet(ws) -> bool:
    """
    Detecta hoja vertical de informe aunque max_column sea alto
    (la plantilla v1.8 pinta columnas grises fuera del formulario).
    """
    a1 = _norm(ws["A1"].value).upper()
    if "INFORME" in a1 and "ID" in a1:
        return True
    # Encabezados de tabla Campo | Valor
    for row in (3, 1, 2, 4):
        a = _norm(ws.cell(row, 1).value).lower()
        b = _norm(ws.cell(row, 2).value).lower()
        if a == "campo" and "valor" in b:
            return True
    return False


def extract_payloads_from_workbook(content: bytes) -> list[dict[str, Any]]:
    """Detecta formato informe (1+ pestañas) o lote y devuelve dicts por caso."""
    payloads: list[dict[str, Any]] = []
    last_names: list[str] = []

    # data_only=True lee valores calculados; si el archivo no se recalculó en Excel,
    # las fórmulas salen vacías. Reintentamos sin data_only.
    for data_only in (True, False):
        try:
            wb = load_workbook(io.BytesIO(content), data_only=data_only)
        except Exception as exc:
            raise ExcelImportError(
                f"No se pudo abrir el Excel ({exc}). "
                "Use un archivo .xlsx de la plantilla descargada del sistema "
                "(no PDF, no .xls antiguo, no capturas)."
            ) from exc
        last_names = list(wb.sheetnames)
        payloads = []

        if "Lote_informes" in wb.sheetnames:
            payloads.extend(parse_lote_sheet(wb["Lote_informes"]))

        for name in wb.sheetnames:
            if name in _SKIP_SHEETS or name.startswith("Catalogo"):
                continue
            ws = wb[name]
            if not (_looks_like_informe_sheet(ws) or _sheet_hab_id_from_name(name) is not None):
                continue
            data = parse_informe_sheet(ws)
            if not data.get("hab_id"):
                hid = _sheet_hab_id_from_name(name)
                if hid is not None:
                    data["hab_id"] = hid
            if data.get("hab_id") or any(
                k in data for k in ("nombre_conf", "fecha_v2", "estado_2da", "decision_D")
            ):
                payloads.append(data)

        if payloads:
            return payloads

    disponibles = ", ".join(last_names) or "(ninguna)"
    raise ExcelImportError(
        "No se encontraron datos de informe ni hoja «Lote_informes». "
        "Use la plantilla Excel descargada desde el sistema "
        f"(pestañas en el archivo: {disponibles})."
    )


def import_excel_bytes(
    content: bytes,
    *,
    caso_esperado: CasoRojo | None = None,
    overwrite_precarga: bool = False,
    user=None,
    solo_visibles_para=None,
) -> dict[str, Any]:
    """
    Aplica el Excel a uno o varios casos.
    Si caso_esperado, exige que el hab_id del archivo coincida (o esté vacío y se asuma).
    Si solo_visibles_para, solo actualiza casos del alcance de ese usuario.
    """
    from inspecciones.workflow import usuario_puede_ver_caso

    payloads = extract_payloads_from_workbook(content)
    results: list[dict[str, Any]] = []
    errors: list[str] = []

    for data in payloads:
        hab_raw = data.get("hab_id")
        try:
            hab_id = int(hab_raw) if hab_raw not in (None, "") else None
        except (TypeError, ValueError):
            errors.append(f"hab_id inválido: {hab_raw!r}")
            continue

        if caso_esperado is not None:
            if hab_id is None:
                hab_id = caso_esperado.hab_id
            elif hab_id != caso_esperado.hab_id:
                errors.append(
                    f"El Excel es del ID Habitable {hab_id}, pero la ficha abierta es {caso_esperado.hab_id}."
                )
                continue
            caso = caso_esperado
        else:
            if hab_id is None:
                errors.append("Fila/hoja sin hab_id.")
                continue
            caso = CasoRojo.objects.filter(hab_id=hab_id).first()
            if not caso:
                errors.append(f"No existe caso con ID Habitable {hab_id}.")
                continue

        if solo_visibles_para is not None and not usuario_puede_ver_caso(solo_visibles_para, caso):
            errors.append(f"ID {caso.hab_id}: no está asignado a su usuario.")
            continue

        try:
            estado_antes = caso.estado_2da
            changed, field_warnings = _apply_dict_to_caso(
                caso, data, overwrite_precarga=overwrite_precarga
            )
            errors.extend(field_warnings)
            caso.sync_gps_texto()
            if changed:
                from django.core.exceptions import ValidationError
                from inspecciones.workflow import TRANSICIONES, validar_transicion_estado

                # Validar transición de estado (mismo grafo/rol que el panel web)
                if "estado_2da" in changed and caso.estado_2da != estado_antes:
                    nuevo = caso.estado_2da
                    try:
                        if user is not None:
                            validar_transicion_estado(user, estado_antes or "", nuevo)
                        else:
                            # Sin usuario: solo grafo + catálogo
                            from inspecciones import choices as ch

                            if nuevo not in {c.value for c in ch.Estado2daRonda}:
                                raise ValidationError(f"Estado no reconocido: «{nuevo}».")
                            if (estado_antes or "") != nuevo and nuevo not in TRANSICIONES.get(
                                estado_antes or "", set()
                            ):
                                raise ValidationError(
                                    f"Transición no permitida: «{estado_antes}» → «{nuevo}». "
                                    f"Siguientes válidos: "
                                    f"{', '.join(sorted(TRANSICIONES.get(estado_antes or '', set())) or ['(ninguno)'])}."
                                )
                    except ValidationError as exc:
                        caso.estado_2da = estado_antes
                        changed = [c for c in changed if c != "estado_2da"]
                        msg = "; ".join(str(m) for m in getattr(exc, "messages", [exc]))
                        errors.append(
                            f"ID {caso.hab_id}: cambio de estado rechazado ({msg}). "
                            "El resto de campos se aplicó si no hay otro error."
                        )
                # Cargar Excel NUNCA exige PDF ni evidencia de cierre:
                # solo vuelca campos. La validación corre al elevar/aprobar.
                if changed:
                    caso.save()
            n_met = _apply_metrados_from_workbook(content, caso)
            if n_met:
                changed = list(changed) + [f"lineas_metrado:{n_met}"]
            n_proc = _apply_procedimientos_sel_from_workbook(content, caso)
            if n_proc:  # None = sin hoja; 0 = sin selección «Sí» → no reportar ruido
                changed = list(changed) + [f"proc_codigos:{n_proc}"]
            results.append(
                {
                    "hab_id": caso.hab_id,
                    "pk": caso.pk,
                    "campos": changed,
                    "nombre": caso.nombre_conf or caso.nombre_hab,
                }
            )
            if changed:
                from inspecciones.historial_detalle import registrar_historial_detallado

                lista = ", ".join(str(c) for c in changed[:25])
                if len(changed) > 25:
                    lista += f" … (+{len(changed) - 25})"
                registrar_historial_detallado(
                    caso=caso,
                    usuario=user,
                    resumen=f"Carga Excel: {len(changed)} cambio(s)",
                    detalle=(
                        f"Se aplicó la plantilla Excel al caso {caso.hab_id}.\n"
                        f"Campos / áreas tocados:\n{lista}"
                    ),
                    origen="excel",
                )
        except ExcelImportError as exc:
            errors.append(f"ID {hab_id}: {exc}")
        except Exception as exc:
            # ValidationError de Django u otros
            msg = getattr(exc, "messages", None)
            if msg:
                errors.append(f"ID {hab_id}: {'; '.join(str(m) for m in msg)}")
            else:
                errors.append(f"ID {hab_id}: {exc}")

    if not results and errors:
        raise ExcelImportError("; ".join(errors))
    if not results:
        raise ExcelImportError(
            "Ningún caso actualizado. Verifique que el Excel sea la plantilla "
            "del sistema y que el ID Habitable coincida con la ficha."
        )

    return {
        "actualizados": results,
        "errores": errors,
        "usuario": getattr(user, "username", None),
    }


def _apply_metrados_from_workbook(content: bytes, caso: CasoRojo) -> int:
    """Si existe hoja Metrados con filas de este hab_id, reemplaza líneas."""
    wb = load_workbook(io.BytesIO(content), data_only=True)
    if "Metrados" not in wb.sheetnames:
        return 0
    ws = wb["Metrados"]
    header_row = None
    headers: list[str] = []
    for i in range(1, 8):
        vals = [_norm(c.value).lower() for c in ws[i]]
        if "hab_id" in vals and "codigo_partida" in vals:
            header_row = i
            headers = vals
            break
    if header_row is None:
        return 0
    idx = {h: i for i, h in enumerate(headers) if h}
    nuevas = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        # Saltar fila de hints didácticos (texto en hab_id)
        try:
            hid = int(row[idx["hab_id"]])
        except (TypeError, ValueError, KeyError):
            continue
        if hid != caso.hab_id:
            continue
        codigo = _norm(row[idx.get("codigo_partida", -1)] if "codigo_partida" in idx else "")
        cantidad_raw = row[idx["cantidad"]] if "cantidad" in idx else None
        # Saltar filas plantilla vacías
        if not codigo and cantidad_raw in (None, ""):
            continue
        # Saltar si codigo_partida es texto de ayuda
        if codigo.lower().startswith("ej."):
            continue
        partida = PartidaCatalogo.objects.filter(codigo__iexact=codigo).first() if codigo else None
        orden = row[idx["orden"]] if "orden" in idx else len(nuevas) + 1
        try:
            orden_i = int(orden) if orden not in (None, "") else len(nuevas) + 1
        except (TypeError, ValueError):
            orden_i = len(nuevas) + 1
        cant = None
        if cantidad_raw not in (None, ""):
            try:
                cant = Decimal(str(cantidad_raw))
            except (InvalidOperation, ValueError):
                cant = None
        nuevas.append(
            LineaMetrado(
                caso=caso,
                orden=orden_i,
                partida=partida,
                codigo_partida=codigo or (partida.codigo if partida else ""),
                elemento=_norm(row[idx["elemento"]]) if "elemento" in idx else "Pendiente",
                id_pln01=_norm(row[idx["id_pln01"]]) if "id_pln01" in idx else "",
                piso_pln=_norm(row[idx["piso_pln"]]) if "piso_pln" in idx else "",
                ubicacion=_norm(row[idx["ubicacion"]]) if "ubicacion" in idx else "",
                tipo_apuntamiento=(
                    _norm(row[idx["tipo_apuntamiento"]]) if "tipo_apuntamiento" in idx else "na"
                )
                or "na",
                cantidad=cant,
                unidad=_norm(row[idx["unidad"]]) if "unidad" in idx else "Pendiente",
                severidad=_norm(row[idx["severidad"]]) if "severidad" in idx else "Pendiente",
                accion=_norm(row[idx["accion"]]) if "accion" in idx else "Pendiente",
                confianza=_norm(row[idx["confianza"]]) if "confianza" in idx else "Pendiente",
                nota=_norm(row[idx["nota"]]) if "nota" in idx else "",
            )
        )
    if not nuevas:
        return 0
    caso.lineas_metrado.all().delete()
    for ln in nuevas:
        # save() alinea codigo/unidad desde la partida del catálogo
        ln.save()
    return len(nuevas)


def _apply_procedimientos_sel_from_workbook(content: bytes, caso: CasoRojo) -> int | None:
    """Si existe Procedimientos_sel, actualiza proc_codigos. None = hoja ausente."""
    wb = load_workbook(io.BytesIO(content), data_only=True)
    sheet_name = None
    for cand in ("Procedimientos_sel", "Procedimientos_Sel", "procedimientos_sel"):
        if cand in wb.sheetnames:
            sheet_name = cand
            break
    if sheet_name is None:
        # coincidencia flexible
        for name in wb.sheetnames:
            n = name.strip().lower().replace(" ", "_")
            if n.startswith("procedimientos") and "sel" in n:
                sheet_name = name
                break
    if sheet_name is None:
        return None
    ws = wb[sheet_name]
    header_row = None
    headers: list[str] = []
    for i in range(1, 8):
        vals = [_norm(c.value).lower() for c in ws[i]]
        # quitar BOM / espacios raros
        vals = [v.lstrip("\ufeff").strip() for v in vals]
        if "codigo" in vals and ("seleccionar" in vals or "seleccion" in vals or "si/no" in vals):
            header_row = i
            headers = vals
            break
    if header_row is None:
        return None
    idx = {h: i for i, h in enumerate(headers) if h}
    sel_key = "seleccionar" if "seleccionar" in idx else (
        "seleccion" if "seleccion" in idx else "si/no"
    )
    codes: list[str] = []
    vio_seleccion = False
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        codigo = _norm(row[idx["codigo"]]).upper()
        sel = _norm(row[idx[sel_key]]).lower()
        if not codigo:
            continue
        if sel in ("sí", "si", "no", "yes", "1", "x", "true", "false", "0"):
            vio_seleccion = True
        if sel in ("sí", "si", "yes", "1", "x", "true"):
            codes.append(codigo)
    if not vio_seleccion:
        return None
    caso.proc_codigos = ", ".join(codes)
    caso.save(update_fields=["proc_codigos", "updated_at"])
    return len(codes)
