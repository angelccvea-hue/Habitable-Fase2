# -*- coding: utf-8 -*-
"""Carga masiva de IDs/certificados (Excel) para asignación por equipo — solo admin."""
from __future__ import annotations

import io
import re
import unicodedata
from typing import Any

from django.contrib.auth.models import User
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from inspecciones.asignacion import etiqueta_equipo, equipos_con_etiqueta
from inspecciones.models import CasoRojo

_HEADERS_HAB = {
    "hab_id",
    "id",
    "id_habitable",
    "idhabitable",
    "habitable_id",
    "nro_id",
    "nr_id",
}
_HEADERS_CERT = {
    "certificado",
    "nr_certificado",
    "nro_certificado",
    "num_certificado",
    "numero_certificado",
    "cert",
    "codigo_certificado",
}
_HEADERS_NOMBRE = {
    "nombre",
    "nombre_edificio",
    "nombre_edificacion",
    "edificio",
    "edificacion",
    "nombre_excel",
    "nombre_habitable",
    "descripcion",
    "inmueble",
}
_HEADERS_ESTADO = {
    "estado",
    "estado_ubicacion",
    "estado_geo",
    "entidad",
    "entidad_federal",
}
_HEADERS_PARROQUIA = {
    "parroquia",
    "parroquia_ubicacion",
}


def unicodedata_fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(ch for ch in s if not unicodedata.combining(ch))


def _norm_header(v: Any) -> str:
    s = str(v or "").strip().lower()
    s = unicodedata_fold(s)
    s = re.sub(r"[^a-z0-9_]+", "_", s)
    return s.strip("_")


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _parse_hab_id(raw: str) -> int | None:
    s = re.sub(r"[^\d]", "", raw or "")
    if not s:
        return None
    try:
        return int(s)
    except ValueError:
        return None


def _parse_cert(raw: str) -> str:
    s = (raw or "").strip()
    if s.startswith('="') and s.endswith('"'):
        s = s[2:-1]
    s = s.strip().strip('"').strip("=")
    s = re.sub(r"\.0$", "", s)
    return s


def _norm_nombre_cmp(s: str) -> str:
    s = unicodedata_fold(str(s or "").lower())
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def comparar_nombres(nombre_excel: str, nombre_hab: str) -> str:
    ex = (nombre_excel or "").strip()
    hab = (nombre_hab or "").strip()
    if not ex:
        return "sin_excel"
    if not hab:
        return "sin_habitable"
    a, b = _norm_nombre_cmp(ex), _norm_nombre_cmp(hab)
    if a == b:
        return "coinciden"
    if a and b and (a in b or b in a):
        return "parcial"
    return "difieren"


def comparar_ubicacion(texto_excel: str, muni_parr_hab: str) -> str:
    """Compara estado/parroquia del Excel contra muni_parr Habitable."""
    ex = (texto_excel or "").strip()
    hab = (muni_parr_hab or "").strip()
    if not ex:
        return "sin_excel"
    if not hab:
        return "sin_habitable"
    a, b = _norm_nombre_cmp(ex), _norm_nombre_cmp(hab)
    if a == b:
        return "coinciden"
    if a and b and (a in b or b in a):
        return "parcial"
    # tokens: si algún token significativo del excel está en hab
    tokens = [t for t in a.split() if len(t) > 3]
    if tokens and any(t in b for t in tokens):
        return "parcial"
    return "difieren"


def plantilla_excel_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "lote"
    headers = [
        "hab_id",
        "certificado",
        "nombre_edificio",
        "estado_ubicacion",
        "parroquia",
    ]
    fill = PatternFill("solid", fgColor="00247D")
    font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    ws.cell(2, 1, 171818)
    ws.cell(2, 2, "")
    ws.cell(2, 3, "Ejemplo edificio A")
    ws.cell(2, 4, "LA GUAIRA")
    ws.cell(2, 5, "Caraballeda")
    ws.cell(3, 1, "")
    ws.cell(3, 2, "73420200720261338")
    ws.cell(3, 3, "Ejemplo por certificado")
    ws.cell(3, 4, "DISTRITO CAPITAL")
    ws.cell(3, 5, "San José")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws2 = wb.create_sheet("instrucciones")
    ws2["A1"] = (
        "Carga masiva Fase II — Asignación por equipo\n\n"
        "OBLIGATORIO por fila: hab_id (ID Habitable) O certificado (al menos uno).\n"
        "Sin ID ni certificado la fila se marca como incompleta y no se asigna.\n\n"
        "Columnas:\n"
        "• hab_id — ID Habitable\n"
        "• certificado — número de certificado\n"
        "• nombre_edificio — opcional; se compara con el nombre Habitable\n"
        "• estado_ubicacion — opcional (entidad federal)\n"
        "• parroquia — opcional; se compara con municipio/parroquia Habitable\n\n"
        "Guarde como .xlsx y cárguelo en Asignación por lote (solo administrador).\n"
        "Revise la vista previa y asigne al equipo (coord.equipoN)."
    )
    ws2["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 95
    ws2.row_dimensions[1].height = 180
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _empty_resumen() -> dict[str, int]:
    return {
        "total": 0,
        "ok": 0,
        "no_encontrado": 0,
        "vacio": 0,
        "duplicado": 0,
        "nombres_coinciden": 0,
        "nombres_parcial": 0,
        "nombres_difieren": 0,
        "nombres_sin_excel": 0,
        "ubic_coinciden": 0,
        "ubic_parcial": 0,
        "ubic_difieren": 0,
        "ya_asignados": 0,
        "libres": 0,
    }


def parsear_excel_lote(content: bytes) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return {
            "filas": [],
            "pks_ok": [],
            "resumen": _empty_resumen(),
            "error": "El Excel está vacío.",
        }

    headers = [_norm_header(h) for h in header_row]
    idx_hab = next((i for i, h in enumerate(headers) if h in _HEADERS_HAB), None)
    idx_cert = next((i for i, h in enumerate(headers) if h in _HEADERS_CERT), None)
    idx_nombre = next((i for i, h in enumerate(headers) if h in _HEADERS_NOMBRE), None)
    idx_estado = next((i for i, h in enumerate(headers) if h in _HEADERS_ESTADO), None)
    idx_parr = next((i for i, h in enumerate(headers) if h in _HEADERS_PARROQUIA), None)
    data_rows = list(rows_iter)
    if idx_hab is None and idx_cert is None:
        idx_hab = 0
        idx_cert = 1 if len(header_row) > 1 else None
        if idx_nombre is None and len(header_row) > 2:
            idx_nombre = 2
        if idx_estado is None and len(header_row) > 3:
            idx_estado = 3
        if idx_parr is None and len(header_row) > 4:
            idx_parr = 4
        data_rows = [header_row] + data_rows

    filas: list[dict] = []
    seen_pks: set[int] = set()
    pks_ok: list[int] = []
    n_ok = n_miss = n_empty = n_dup = 0
    n_coin = n_parc = n_dif = n_sin_ex = 0
    u_coin = u_parc = u_dif = 0
    n_ya = n_libres = 0

    def _col(idx: int | None) -> str:
        if idx is None or idx >= len(vals):
            return ""
        return _cell_str(vals[idx])

    for n, row in enumerate(data_rows, start=2):
        if row is None:
            continue
        vals = list(row)
        hab_raw = _col(idx_hab)
        cert_raw = _parse_cert(_col(idx_cert)) if idx_cert is not None else ""
        nombre_excel = _col(idx_nombre)
        estado_excel = _col(idx_estado)
        parroquia_excel = _col(idx_parr)

        if not any(_cell_str(v) for v in vals):
            continue

        hab_id = _parse_hab_id(hab_raw) if hab_raw else None
        # Obligatorio: hab_id O certificado
        if not hab_id and not cert_raw:
            n_empty += 1
            filas.append(
                {
                    "fila": n,
                    "hab_id": "",
                    "certificado": "",
                    "nombre_excel": nombre_excel,
                    "nombre_habitable": "",
                    "nombre_match": "sin_excel" if not nombre_excel else "sin_habitable",
                    "estado_excel": estado_excel,
                    "parroquia_excel": parroquia_excel,
                    "ubicacion_habitable": "",
                    "ubicacion_match": "sin_excel",
                    "estado": "vacio",
                    "caso_pk": None,
                    "nombre": "",
                    "equipo_actual": "",
                    "detalle": "Obligatorio: indique hab_id o certificado",
                }
            )
            continue

        caso = None
        if hab_id:
            caso = CasoRojo.objects.filter(hab_id=hab_id).first()
        if caso is None and cert_raw:
            caso = CasoRojo.objects.filter(certificado__iexact=cert_raw).first()

        if caso is None:
            n_miss += 1
            filas.append(
                {
                    "fila": n,
                    "hab_id": hab_id or "",
                    "certificado": cert_raw,
                    "nombre_excel": nombre_excel,
                    "nombre_habitable": "",
                    "nombre_match": "sin_habitable",
                    "estado_excel": estado_excel,
                    "parroquia_excel": parroquia_excel,
                    "ubicacion_habitable": "",
                    "ubicacion_match": "sin_habitable",
                    "estado": "no_encontrado",
                    "caso_pk": None,
                    "nombre": "",
                    "equipo_actual": "",
                    "detalle": "No hay caso con ese ID/certificado",
                }
            )
            continue

        nombre_hab = (caso.nombre_conf or caso.nombre_hab or "").strip()
        muni_parr = (caso.muni_parr or "").strip()
        match = comparar_nombres(nombre_excel, nombre_hab)
        # Validar ubicación: priorizar parroquia; si no, estado
        ubic_txt = parroquia_excel or estado_excel
        ubic_match = comparar_ubicacion(ubic_txt, muni_parr)
        if match == "coinciden":
            n_coin += 1
        elif match == "parcial":
            n_parc += 1
        elif match == "difieren":
            n_dif += 1
        else:
            n_sin_ex += 1
        if ubic_match == "coinciden":
            u_coin += 1
        elif ubic_match == "parcial":
            u_parc += 1
        elif ubic_match == "difieren":
            u_dif += 1

        if caso.pk in seen_pks:
            n_dup += 1
            estado = "duplicado"
            detalle = "Ya listado en una fila anterior"
            ya_asignado = bool(caso.coordinador_asignado_id)
        else:
            seen_pks.add(caso.pk)
            pks_ok.append(caso.pk)
            n_ok += 1
            ya_asignado = bool(caso.coordinador_asignado_id)
            if ya_asignado:
                n_ya += 1
                estado = "ya_asignado"
                detalle = "Ya tiene equipo — no se reasignará por defecto"
            else:
                n_libres += 1
                estado = "ok"
                detalle = "Listo para asignar (sin coordinador)"
            if match == "difieren":
                detalle += " · nombre Excel ≠ Habitable"
            elif match == "parcial":
                detalle += " · nombre parcial"
            if ubic_match == "difieren":
                detalle += " · ubicación difiere"

        coord = caso.coordinador_asignado
        equipo_act = ""
        if coord:
            from inspecciones.asignacion import equipo_de_username

            eq = equipo_de_username(coord.username or "")
            equipo_act = etiqueta_equipo(eq) if eq else (coord.username or "")

        filas.append(
            {
                "fila": n,
                "hab_id": caso.hab_id,
                "certificado": caso.certificado or cert_raw,
                "nombre_excel": nombre_excel,
                "nombre_habitable": nombre_hab[:120],
                "nombre_match": match,
                "estado_excel": estado_excel,
                "parroquia_excel": parroquia_excel,
                "ubicacion_habitable": muni_parr[:120],
                "ubicacion_match": ubic_match,
                "estado": estado,
                "ya_asignado": ya_asignado,
                "caso_pk": caso.pk,
                "nombre": nombre_hab[:80],
                "banda": caso.banda or "",
                "estado_2da": caso.estado_2da or "",
                "score": caso.score,
                "equipo_actual": equipo_act,
                "detalle": detalle,
            }
        )

    return {
        "filas": filas,
        "pks_ok": pks_ok,
        "resumen": {
            "total": len(filas),
            "ok": n_ok,
            "no_encontrado": n_miss,
            "vacio": n_empty,
            "duplicado": n_dup,
            "nombres_coinciden": n_coin,
            "nombres_parcial": n_parc,
            "nombres_difieren": n_dif,
            "nombres_sin_excel": n_sin_ex,
            "ubic_coinciden": u_coin,
            "ubic_parcial": u_parc,
            "ubic_difieren": u_dif,
            "ya_asignados": n_ya,
            "libres": n_libres,
        },
        "error": None,
    }


def opciones_equipos_asignacion() -> list[dict[str, Any]]:
    out = []
    etiquetas = {e["numero"]: e for e in equipos_con_etiqueta()}
    coords = (
        User.objects.filter(username__istartswith="coord.equipo", is_active=True, is_staff=True)
        .order_by("username")
    )
    for u in coords:
        m = re.match(r"^coord\.equipo(\d+)$", u.username, re.I)
        if not m:
            continue
        num = m.group(1)
        et = etiquetas.get(num, {})
        out.append(
            {
                "numero": num,
                "etiqueta": et.get("etiqueta") or etiqueta_equipo(num),
                "coordinador_id": u.pk,
                "coordinador_username": u.username,
                "coordinador_nombre": (u.get_full_name() or u.username).strip(),
            }
        )
    return out
